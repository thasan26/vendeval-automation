"""Weekly outputs: an HTML status dashboard, an Excel workbook, and a data-quality report.

The HTML dashboard is the "no license needed" version of the Power BI report
(powerbi/ has the measures and page layout to build the same thing in Power BI).
"""
from __future__ import annotations

import html
from datetime import date, timedelta

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import REPORTS, STATUS_ORDER

# Sequential blue ramp (lifecycle is ordered, so it is a magnitude, not identities)
STATUS_FILL = {"Not Started": "#86b6ef", "Requested": "#5598e7", "In Progress": "#2a78d6",
               "Submitted": "#1c5cab", "Approved": "#0d366b"}
TIER_LABEL = {0: "Not overdue", 1: "1-29 days", 2: "30-89 days", 3: "90-179 days", 4: "180+ days"}

esc = html.escape


# ------------------------------------------------------------------ helpers
def weekly_throughput(con, as_of: date, weeks: int = 12) -> pd.DataFrame:
    h = pd.read_sql_query("SELECT status, changed_at FROM status_history", con, parse_dates=["changed_at"])
    rows = []
    for k in range(weeks - 1, -1, -1):
        end = pd.Timestamp(as_of - timedelta(days=7 * k)) + pd.Timedelta(days=1)
        start = end - pd.Timedelta(days=7)
        wk = h[(h.changed_at >= start) & (h.changed_at < end)]
        rows.append({"week_ending": (end - pd.Timedelta(days=1)).date(),
                     "requested": (wk.status == "Requested").sum(),
                     "submitted": (wk.status == "Submitted").sum(),
                     "approved": (wk.status == "Approved").sum()})
    return pd.DataFrame(rows)


def svg_stacked_bars(df: pd.DataFrame, label_col: str, cats: list[str], colors: dict) -> str:
    """Horizontal 100% stacked bars with a 2px surface gap between segments."""
    w, row_h, bar_h, left = 1100, 40, 26, 50
    plot_w = w - left - 60
    out = [f'<svg viewBox="0 0 {w} {row_h * len(df) + 10}" role="img" class="chart">']
    for i, (_, r) in enumerate(df.iterrows()):
        total = sum(r[c] for c in cats)
        y = 6 + i * row_h
        out.append(f'<text x="{left - 10}" y="{y + bar_h / 2 + 5}" text-anchor="end" class="axis-label">{r[label_col]}</text>')
        x = left
        for c in cats:
            seg = plot_w * r[c] / total if total else 0
            if seg <= 0:
                continue
            pct = 100 * r[c] / total
            out.append(f'<rect class="mark" x="{x:.1f}" y="{y}" width="{max(seg - 2, 1):.1f}" height="{bar_h}" rx="3" '
                       f'fill="{colors[c]}" data-tip="{r[label_col]} · {c}: {r[c]:,} ({pct:.0f}%)"></rect>')
            if seg > 44:
                out.append(f'<text x="{x + seg / 2 - 1:.1f}" y="{y + bar_h / 2 + 4}" text-anchor="middle" class="in-bar">{pct:.0f}%</text>')
            x += seg
        out.append(f'<text x="{left + plot_w + 8}" y="{y + bar_h / 2 + 5}" class="value">{total:,}</text>')
    out.append("</svg>")
    return "".join(out)


def svg_column_chart(labels: list[str], values: list[int], tip_fmt: str, color="var(--series-1)") -> str:
    w, h, left, bottom, top = 640, 220, 40, 30, 14
    plot_w, plot_h = w - left - 10, h - bottom - top
    vmax = max(max(values), 1)
    step = 10 ** max(len(str(int(vmax))) - 1, 0)
    nice = ((vmax // step) + 1) * step
    bw = plot_w / len(values)
    out = [f'<svg viewBox="0 0 {w} {h}" role="img" class="chart">']
    for t in range(0, 5):
        val = nice * t / 4
        y = top + plot_h - plot_h * val / nice
        out.append(f'<line x1="{left}" x2="{w - 10}" y1="{y:.1f}" y2="{y:.1f}" class="grid"></line>'
                   f'<text x="{left - 6}" y="{y + 4:.1f}" text-anchor="end" class="axis-label">{val:,.0f}</text>')
    for i, (lab, v) in enumerate(zip(labels, values)):
        bh = plot_h * v / nice
        x = left + i * bw + bw * 0.18
        out.append(f'<path class="mark" fill="{color}" data-tip="{tip_fmt.format(label=lab, value=v)}" '
                   f'd="M{x:.1f},{top + plot_h:.1f} v{-max(bh - 4, 0):.1f} q0,-4 4,-4 h{bw * 0.64 - 8:.1f} '
                   f'q4,0 4,4 v{max(bh - 4, 0):.1f} z"></path>')
        if i % 2 == len(labels) % 2 or len(labels) <= 6:
            out.append(f'<text x="{x + bw * 0.32:.1f}" y="{h - 10}" text-anchor="middle" class="axis-label">{lab}</text>')
    out.append(f'<line x1="{left}" x2="{w - 10}" y1="{top + plot_h}" y2="{top + plot_h}" class="baseline"></line></svg>')
    return "".join(out)


def table_html(df: pd.DataFrame, num_cols=(), cls="") -> str:
    head = "".join(f"<th class='{'num' if c in num_cols else ''}'>{esc(str(c))}</th>" for c in df.columns)
    body = "".join("<tr>" + "".join(
        f"<td class='{'num' if c in num_cols else ''}'>{esc(f'{v:,}' if isinstance(v, int) else str(v))}</td>"
        for c, v in zip(df.columns, row)) + "</tr>" for row in df.itertuples(index=False))
    return f"<div class='table-wrap'><table class='{cls}'><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table></div>"


# ------------------------------------------------------------------ dashboard
def build_dashboard(con, out, checks, as_of: date) -> str:
    ev = pd.read_sql_query("SELECT * FROM vw_evaluation_detail", con)
    wq = out["work_queue"]
    total = len(ev)
    approved = (ev.status == "Approved").sum()
    pm_open = wq[wq.status != "Submitted"]
    overdue = (pm_open.days_overdue > 0).sum()
    over180 = (wq.escalation_tier == 4).sum()
    bidding_vendors = wq[wq.vendor_bidding].vendor_id.nunique()
    thr = weekly_throughput(con, as_of)
    this_wk, last_wk = int(thr.approved.iloc[-1]), int(thr.approved.iloc[-2])
    delta = this_wk - last_wk
    all_pass = bool(checks.passed.all())

    tiles = [
        ("Evaluations tracked", f"{total:,}", "2022–2024 contracts"),
        ("Approved", f"{100 * approved / total:.0f}%", f"{approved:,} of {total:,}"),
        ("Waiting on PMs", f"{len(pm_open):,}", f"{overdue:,} overdue"),
        ("180+ days overdue", f"{over180:,}", "at leadership escalation"),
        ("Approved this week", f"{this_wk}",
         f"{'▲' if delta > 0 else '▼' if delta < 0 else '■'} {abs(delta)} vs last week"),
        ("Vendors bidding now", f"{bidding_vendors}", "with open evaluations"),
    ]
    tiles_html = "".join(f"<div class='tile'><div class='tile-label'>{a}</div><div class='tile-value'>{b}</div>"
                         f"<div class='tile-sub'>{c}</div></div>" for a, b, c in tiles)

    by_year = ev.pivot_table(index="contract_year", columns="status", values="eval_id", aggfunc="count", fill_value=0)
    by_year = by_year.reindex(columns=STATUS_ORDER, fill_value=0).reset_index()
    by_year["contract_year"] = by_year.contract_year.astype(str)
    legend = "".join(f"<span class='key'><i style='background:{STATUS_FILL[s]}'></i>{s}</span>" for s in STATUS_ORDER)
    year_chart = svg_stacked_bars(by_year, "contract_year", STATUS_ORDER, STATUS_FILL)

    thr_chart = svg_column_chart([f"{d:%m/%d}" for d in thr.week_ending], thr.approved.tolist(),
                                 "Week ending {label}: {value} approved")
    tiers = wq[wq.status != "Submitted"].escalation_tier.value_counts().reindex(range(5), fill_value=0)
    tier_chart = svg_column_chart([TIER_LABEL[t] for t in tiers.index], tiers.tolist(),
                                  "{label}: {value} evaluations")

    dept = (wq.merge(ev[["eval_id"]], on="eval_id")
              .groupby("department").agg(open=("eval_id", "count"),
                                         overdue=("days_overdue", lambda s: int((s > 0).sum())),
                                         over_180=("escalation_tier", lambda s: int((s == 4).sum())),
                                         stalled=("stalled", "sum"),
                                         avg_days_overdue=("days_overdue", lambda s: int(s[s > 0].mean())))
              .sort_values("over_180", ascending=False).reset_index())
    dept["stalled"] = dept.stalled.astype(int)
    dept.columns = ["Department", "Open", "Overdue", "180+ days", "Stalled", "Avg days overdue"]

    vend = ev.set_index("eval_id").vendor_name
    top = wq.nsmallest(15, "priority_rank").copy()
    top["Vendor"] = top.eval_id.map(vend)
    names = out["evaluators"].set_index("evaluator_id")
    top["Owner"] = top.assigned_to.map(lambda e: f"{names.at[e, 'first_name']} {names.at[e, 'last_name']}")
    top["Flags"] = [" · ".join(f for f, on in [("Vendor bidding", b), ("Stalled", s), ("Reassigned", "reassigned" in r)] if on)
                    for b, s, r in zip(top.vendor_bidding, top.stalled, top.assignment_reason)]
    top = top[["priority_rank", "eval_id", "Vendor", "status", "days_overdue", "Owner", "escalation_tier", "Flags"]]
    top.columns = ["#", "Evaluation", "Vendor", "Status", "Days overdue", "Owner", "Tier", "Flags"]
    top["Days overdue"] = top["Days overdue"].astype(int)

    chk = checks.copy()
    chk["Result"] = chk.passed.map({True: "✓ Pass", False: "✕ Fail"})
    chk = chk[["Result", "check", "detail"]].rename(columns={"check": "Check", "detail": "Detail"})
    thr_table = thr.rename(columns={"week_ending": "Week ending", "requested": "Requested",
                                    "submitted": "Submitted", "approved": "Approved"})

    banner = ("<div class='banner ok'>✓ All data-quality checks passed — report cleared to publish</div>" if all_pass else
              "<div class='banner bad'>✕ Data-quality checks failed — do not publish (see checks table)</div>")
    return f"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>VENDEVAL Weekly Status</title>
<style>
:root{{--page:#f9f9f7;--surface:#fcfcfb;--ink:#0b0b0b;--ink2:#52514e;--muted:#898781;--grid:#e1e0d9;
--base:#c3c2b7;--series-1:#2a78d6;--good:#006300;--crit:#d03b3b;--border:#e1e0d9}}
@media (prefers-color-scheme:dark){{:root:not([data-theme="light"]){{--page:#0d0d0d;--surface:#1a1a19;--ink:#f0efec;
--ink2:#c3c2b7;--grid:#2c2c2a;--base:#383835;--series-1:#3987e5;--good:#0ca30c;--border:#2c2c2a}}}}
*{{box-sizing:border-box}} body{{margin:0;background:var(--page);color:var(--ink);
font:14px/1.45 system-ui,-apple-system,"Segoe UI",sans-serif;font-variant-numeric:tabular-nums}}
main{{max-width:1180px;margin:0 auto;padding:24px 16px 48px}}
h1{{font-size:22px;margin:0}} h2{{font-size:15px;margin:0 0 4px}} .sub{{color:var(--ink2);margin:4px 0 16px}}
.note{{color:var(--ink2);font-size:12.5px;margin:0 0 10px}}
.banner{{padding:10px 14px;border-radius:8px;margin:0 0 16px;font-weight:600;border:1px solid var(--border);background:var(--surface)}}
.banner.ok{{color:var(--good)}} .banner.bad{{color:var(--crit)}}
.tiles{{display:grid;grid-template-columns:repeat(auto-fit,minmax(160px,1fr));gap:12px;margin-bottom:16px}}
.tile,.card{{background:var(--surface);border:1px solid var(--border);border-radius:10px;padding:14px 16px}}
.tile-label{{color:var(--ink2);font-size:12.5px}} .tile-value{{font-size:28px;font-weight:650;margin:2px 0}}
.tile-sub{{color:var(--muted);font-size:12.5px}}
.grid2{{display:grid;grid-template-columns:repeat(auto-fit,minmax(340px,1fr));gap:16px;margin-bottom:16px}}
.card{{margin-bottom:16px}} .grid2 .card{{margin-bottom:0}}
.chart{{width:100%;height:auto;display:block}} .chart text{{font-size:12px;fill:var(--ink2)}}
.chart .value{{fill:var(--ink);font-weight:600}} .chart .in-bar{{fill:#fff;font-size:11px}}
.grid{{stroke:var(--grid);stroke-width:1}} .baseline{{stroke:var(--base);stroke-width:1}}
.mark{{cursor:default}} .mark:hover{{opacity:.8}}
.legend{{display:flex;flex-wrap:wrap;gap:6px 14px;margin:6px 0 8px;font-size:12.5px;color:var(--ink2)}}
.key i{{display:inline-block;width:10px;height:10px;border-radius:2px;margin-right:6px;vertical-align:-1px}}
.table-wrap{{overflow-x:auto}} table{{border-collapse:collapse;width:100%;font-size:13px}}
th{{text-align:left;color:var(--ink2);font-weight:600;border-bottom:1px solid var(--base);padding:6px 8px;white-space:nowrap}}
td{{border-bottom:1px solid var(--grid);padding:6px 8px}} td:first-child{{white-space:nowrap}} .num{{text-align:right}}
details summary{{cursor:pointer;color:var(--ink2);font-size:12.5px;margin-top:8px}}
#tip{{position:fixed;pointer-events:none;background:var(--ink);color:var(--surface);padding:5px 8px;border-radius:6px;
font-size:12px;opacity:0;transition:opacity .1s;z-index:9}}
</style></head><body><main>
<h1>VENDEVAL Weekly Status — Supplier Relations</h1>
<p class="sub">Week ending {as_of:%A, %B %d, %Y} · merged from Access, Excel tracker and Archer · synthetic data</p>
{banner}
<div class="tiles">{tiles_html}</div>
<div class="card"><h2>Evaluation status by contract year</h2>
<p class="note">Share of each year's evaluations at each stage. Older years should be closest to fully approved.</p>
<div class="legend">{legend}</div><div class="table-wrap"><div style="min-width:640px">{year_chart}</div></div>
<details><summary>Show as table</summary>{table_html(by_year.rename(columns={"contract_year": "Contract year"}), STATUS_ORDER)}</details></div>
<div class="grid2">
<div class="card"><h2>Evaluations approved per week</h2><p class="note">Last 12 weeks. The pace needed to clear the backlog is the target.</p>{thr_chart}
<details><summary>Show as table</summary>{table_html(thr_table, ["Requested", "Submitted", "Approved"])}</details></div>
<div class="card"><h2>Outstanding evaluations by days overdue</h2><p class="note">Waiting on project managers. Each band maps to an escalation tier.</p>{tier_chart}</div>
</div>
<div class="card"><h2>Outreach priority — top 15</h2><p class="note">Ranked by days overdue, contract value, contract year and whether the vendor is bidding on new MTA work.</p>
{table_html(top, ["#", "Days overdue", "Tier"])}</div>
<div class="card"><h2>Backlog by department</h2>{table_html(dept, ["Open", "Overdue", "180+ days", "Stalled", "Avg days overdue"])}</div>
<div class="card"><h2>Data-quality and reporting-integrity checks</h2>{table_html(chk)}</div>
</main><div id="tip"></div>
<script>
const tip=document.getElementById('tip');
document.querySelectorAll('[data-tip]').forEach(el=>{{
 el.addEventListener('mousemove',e=>{{tip.textContent=el.dataset.tip;tip.style.opacity=1;
  tip.style.left=Math.min(e.clientX+12,innerWidth-tip.offsetWidth-8)+'px';tip.style.top=(e.clientY-34)+'px'}});
 el.addEventListener('mouseleave',()=>tip.style.opacity=0)}});
</script></body></html>"""


# ------------------------------------------------------------------ excel
def build_workbook(con, out, checks, as_of: date, path):
    ev = pd.read_sql_query("SELECT * FROM vw_evaluation_detail", con)
    wq = out["work_queue"].merge(ev[["eval_id", "vendor_name", "pm_name"]], on="eval_id", how="left")
    names = out["evaluators"].set_index("evaluator_id")
    wq["assigned_name"] = wq.assigned_to.map(lambda e: f"{names.at[e, 'first_name']} {names.at[e, 'last_name']}")
    queue = wq[["priority_rank", "eval_id", "contract_no", "contract_year", "department", "vendor_name", "status",
                "due_date", "days_overdue", "sla_state", "escalation_tier", "stalled", "vendor_bidding",
                "pm_name", "assigned_name", "assignment_reason", "priority_score"]].sort_values("priority_rank")
    summary = pd.read_sql_query("SELECT * FROM vw_backlog_by_year", con)
    workload = pd.read_sql_query("SELECT * FROM vw_evaluator_workload ORDER BY open_evals DESC", con)
    vendor_perf = pd.read_sql_query("SELECT * FROM vw_vendor_performance ORDER BY avg_score", con)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in [("Summary", summary), ("Work Queue", queue), ("Evaluator Workload", workload),
                         ("Notifications", out["notifications"].drop(columns=["body"])),
                         ("Vendor Performance", vendor_perf), ("Data Issues", out["data_issues"]),
                         ("QA Checks", checks)]:
            df.to_excel(xw, sheet_name=name, index=False)
    wb = load_workbook(path)
    head_fill = PatternFill("solid", fgColor="1C5CAB")
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), head_fill
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for i, col in enumerate(ws.columns, 1):
            width = max(len(str(c.value or "")) for c in list(col)[:200])
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 48)
    ws = wb["Work Queue"]
    col = get_column_letter(list(queue.columns).index("days_overdue") + 1)
    rng = f"{col}2:{col}{ws.max_row}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["180"],
                                                  fill=PatternFill("solid", fgColor="F4C7C7")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="between", formula=["30", "179"],
                                                  fill=PatternFill("solid", fgColor="FDE7B0")))
    wb["Summary"]["H1"] = f"Report date: {as_of:%Y-%m-%d}"
    wb.save(path)


# ------------------------------------------------------------------ dq report
def build_quality_report(out, checks, as_of) -> str:
    st, issues = out["stats"], out["data_issues"]
    by = issues.assign(issue=issues.issue.str.replace(r"\(score [0-9.]+\)", "", regex=True)
                       .str.replace(r"sources disagree: .*", "sources disagree on status", regex=True)) \
        .groupby(["source", "issue"]).size().reset_index(name="count").sort_values("count", ascending=False)
    lines = [f"# Data Quality Report — {as_of:%Y-%m-%d}", "",
             "## Record flow", "",
             "| Stage | Records |", "|---|---:|"]
    for s, n in st["source_rows"].items():
        lines.append(f"| Source rows: {s} | {n:,} |")
    lines += [f"| After validation (quarantined rows removed) | {st['records_after_validation']:,} |",
              f"| After within-source de-duplication | {st['records_after_dedupe']:,} |",
              f"| **Golden evaluation records** | **{st['golden_records']:,}** |", "",
              "## Vendor name matching", "", "| Method | Records |", "|---|---:|"]
    for m, n in sorted(st["vendor_match_methods"].items(), key=lambda kv: -kv[1]):
        lines.append(f"| {m} | {n:,} |")
    lines += ["", "## Issues found and fixed", "", "| Source | Issue | Count |", "|---|---|---:|"]
    lines += [f"| {r.source} | {r.issue.strip()} | {r.count:,} |" for r in by.itertuples()]
    lines += ["", "## Checks", "", "| Result | Check | Detail |", "|---|---|---|"]
    lines += [f"| {'PASS' if r.passed else 'FAIL'} | {r.check} | {r.detail} |" for r in checks.itertuples()]
    return "\n".join(lines) + "\n"


def write_all(con, out, checks, as_of: date):
    REPORTS.mkdir(exist_ok=True)
    (REPORTS / "weekly_dashboard.html").write_text(build_dashboard(con, out, checks, as_of), encoding="utf-8")
    build_workbook(con, out, checks, as_of, REPORTS / "VENDEVAL_Weekly_Report.xlsx")
    (REPORTS / "data_quality_report.md").write_text(build_quality_report(out, checks, as_of), encoding="utf-8")
